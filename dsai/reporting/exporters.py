"""Export an analysis: Markdown, HTML, Excel, JSON, PDF and Python."""

from __future__ import annotations

import html as _html
import io as _io
import json
from pathlib import Path
from typing import Any

import pandas as pd

from dsai.engines.metrics import human_number
from dsai.reporting.builder import Report, build_report
from dsai.reporting.figures import to_png

_HTML_TEMPLATE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
  :root {{ --ink:#1a1d21; --muted:#5c6470; --line:#e3e6ea; --bg:#fdfdfc; --accent:#2b6cb0;
           --warn:#b7791f; --crit:#c53030; }}
  @media (prefers-color-scheme: dark) {{
    :root {{ --ink:#e8eaed; --muted:#9aa3af; --line:#2f343b; --bg:#16181c; --accent:#7fb3e8;
             --warn:#e0b357; --crit:#f08a8a; }}
  }}
  body {{ margin:0; background:var(--bg); color:var(--ink);
          font:16px/1.65 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif; }}
  .wrap {{ max-width:52rem; margin:0 auto; padding:3rem 1.5rem 6rem; }}
  h1 {{ font-size:2rem; line-height:1.2; margin:0 0 .25rem; letter-spacing:-.02em; }}
  h2 {{ font-size:1.3rem; margin:2.75rem 0 .75rem; padding-bottom:.4rem;
        border-bottom:1px solid var(--line); letter-spacing:-.01em; }}
  h3 {{ font-size:1.05rem; margin:1.75rem 0 .5rem; color:var(--muted);
        text-transform:uppercase; letter-spacing:.06em; font-weight:600; }}
  .meta {{ color:var(--muted); font-size:.875rem; margin-bottom:2rem; }}
  table {{ border-collapse:collapse; width:100%; margin:1rem 0; font-size:.9rem; }}
  th,td {{ border-bottom:1px solid var(--line); padding:.5rem .6rem; text-align:left; vertical-align:top; }}
  th {{ font-weight:600; color:var(--muted); font-size:.8rem;
        text-transform:uppercase; letter-spacing:.04em; }}
  code {{ background:color-mix(in srgb, var(--line) 55%, transparent); padding:.1em .35em;
          border-radius:3px; font-size:.9em; }}
  pre {{ background:color-mix(in srgb, var(--line) 40%, transparent); padding:1rem;
         border-radius:6px; overflow-x:auto; font-size:.82rem; line-height:1.5; }}
  ul {{ padding-left:1.25rem; }}
  li {{ margin:.3rem 0; }}
  .scroll {{ overflow-x:auto; }}
  blockquote {{ margin:1rem 0; padding:.6rem 1rem; border-left:3px solid var(--accent);
                color:var(--muted); }}
  em {{ color:var(--muted); }}
  figure {{ margin:1.5rem 0 2rem; }}
  figure h3 {{ margin:0 0 .5rem; }}
  figcaption {{ color:var(--muted); font-size:.875rem; line-height:1.55; margin-top:.6rem; }}
  figure img {{ max-width:100%; height:auto; display:block; }}
  details.data {{ margin-top:.75rem; border-top:1px solid var(--line); padding-top:.5rem; }}
  details.data summary {{ cursor:pointer; font-size:.85rem; color:var(--accent); }}
  details.data[open] summary {{ margin-bottom:.5rem; }}
  a:focus-visible, summary:focus-visible {{ outline:2px solid var(--accent); outline-offset:2px; }}
  @media print {{
    details.data {{ display:block; }}
    details.data > *:not(summary) {{ display:block !important; }}
  }}
</style></head>
<body><div class="wrap">{body}</div></body></html>
"""


def to_markdown(
    run: Any,
    audience: str = "both",
    frame: pd.DataFrame | None = None,
    chart_files: dict[str, str] | None = None,
    methodology: bool = False,
) -> str:
    report = build_report(run, audience, frame=frame, include_charts=True,
                          include_methodology=methodology)
    return report.to_markdown(chart_files=chart_files)


def to_html(run: Any, audience: str = "both", frame: pd.DataFrame | None = None,
            charts: bool = True, methodology: bool = False) -> str:
    """A single self-contained HTML file: text, tables and interactive charts.

    The Plotly runtime is inlined into the first chart rather than fetched from a
    CDN, so the report opens on a machine with no internet — which is often the
    machine a report is read on. That costs about 3 MB once, not per chart.
    """
    report = build_report(run, audience, frame=frame, include_charts=charts,
                          include_methodology=methodology)
    body = [f"<h1>{_html.escape(report.title)}</h1>",
            f'<p class="meta">Generated {_html.escape(report.generated_at)}</p>']
    if report.executive_summary:
        body.append("<h2>Executive summary</h2>")
        body.append(_markdown_to_html(report.executive_summary))

    first_chart = True
    for heading, content in report.sections:
        body.append(f"<h2>{_html.escape(heading)}</h2>")
        body.append(_markdown_to_html(content))
        for figure in report.figures_for(heading):
            rendered, first_chart = _figure_html(figure, include_js=first_chart)
            if rendered:
                body.append(rendered)
    return _HTML_TEMPLATE.format(title=_html.escape(report.title), body="\n".join(body))


def _figure_html(figure: Any, include_js: bool) -> tuple[str, bool]:
    """One ``<figure>``: the chart, its caption, and the numbers behind it.

    The table is always written, inside a collapsed ``<details>``. A chart that
    fails to render leaves the caption and the table, which still say what the
    chart was there to say.
    """
    from dsai.reporting.figures import to_html_div

    div = to_html_div(figure.figure, include_js=include_js)
    parts = [f"<figure><h3>{_html.escape(figure.title)}</h3>"]
    if div:
        parts.append(f'<div role="img" aria-label="{_html.escape(figure.alt)}">{div}</div>')
        include_js = False
    parts.append(f"<figcaption>{_html.escape(figure.caption)}</figcaption>")
    if figure.has_table:
        parts.append(
            '<details class="data"><summary>The numbers behind this chart</summary>'
            f'<div class="scroll">{_frame_to_html(figure.table)}</div></details>'
        )
    parts.append("</figure>")
    return "".join(parts), include_js


def _frame_to_html(frame: pd.DataFrame) -> str:
    header = "".join(f"<th scope=\"col\">{_html.escape(str(c))}</th>" for c in frame.columns)
    rows = []
    for row in frame.head(40).itertuples(index=False):
        cells = "".join(f"<td>{_html.escape('' if v is None else str(v))}</td>" for v in row)
        rows.append(f"<tr>{cells}</tr>")
    return f"<table><thead><tr>{header}</tr></thead><tbody>{''.join(rows)}</tbody></table>"


def _markdown_to_html(text: str) -> str:
    """A small markdown subset: enough for these reports, no dependency."""
    import re

    out: list[str] = []
    lines = text.split("\n")
    i = 0
    in_list = False
    in_code = False

    def close_list() -> None:
        nonlocal in_list
        if in_list:
            out.append("</ul>")
            in_list = False

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            close_list()
            if in_code:
                out.append("</pre>")
                in_code = False
            else:
                out.append("<pre>")
                in_code = True
            i += 1
            continue
        if in_code:
            out.append(_html.escape(line))
            i += 1
            continue

        # markdown table
        if line.strip().startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|- :"):
            close_list()
            headers = [c.strip() for c in line.strip().strip("|").split("|")]
            out.append('<div class="scroll"><table><thead><tr>'
                       + "".join(f"<th>{_inline(h)}</th>" for h in headers)
                       + "</tr></thead><tbody>")
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                out.append("<tr>" + "".join(f"<td>{_inline(c)}</td>" for c in cells) + "</tr>")
                i += 1
            out.append("</tbody></table></div>")
            continue

        stripped = line.strip()
        if stripped.startswith("### "):
            close_list()
            out.append(f"<h3>{_inline(stripped[4:])}</h3>")
        elif stripped.startswith("## "):
            close_list()
            out.append(f"<h3>{_inline(stripped[3:])}</h3>")
        elif re.match(r"^[-*] ", stripped):
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{_inline(stripped[2:])}</li>")
        elif re.match(r"^\d+\. ", stripped):
            if not in_list:
                out.append("<ul>")
                in_list = True
            item = re.sub(r"^\d+\. ", "", stripped)
            out.append(f"<li>{_inline(item)}</li>")
        elif not stripped:
            close_list()
        else:
            close_list()
            out.append(f"<p>{_inline(stripped)}</p>")
        i += 1
    close_list()
    if in_code:
        out.append("</pre>")
    return "\n".join(out)


def _inline(text: str) -> str:
    import re

    escaped = _html.escape(text)
    escaped = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", escaped)
    escaped = re.sub(r"(?<!\*)\*([^*]+?)\*(?!\*)", r"<em>\1</em>", escaped)
    escaped = re.sub(r"`([^`]+?)`", r"<code>\1</code>", escaped)
    return escaped


def to_excel(run: Any, path: str | Path, frame: pd.DataFrame | None = None) -> Path:
    """Multi-sheet workbook: summary, tournament, findings, recommendations, profile."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    from dsai.repro.provenance import build_manifest

    manifest = build_manifest(run)
    sheets: dict[str, pd.DataFrame] = {}

    sheets["Summary"] = pd.DataFrame(
        [
            {"Item": "Dataset", "Value": run.dataset_name},
            {"Item": "Rows", "Value": run.profile.n_rows if run.profile else None},
            {"Item": "Columns", "Value": run.profile.n_columns if run.profile else None},
            {"Item": "Data quality score", "Value": run.profile.quality_score if run.profile else None},
            {"Item": "Task", "Value": run.objective.task_type.value if run.objective else None},
            {"Item": "Target", "Value": run.objective.target if run.objective else None},
            {"Item": "Selected model", "Value": run.best.model_name if run.best else None},
            {"Item": "Primary metric", "Value": run.plan.primary_metric if run.plan else None},
            {"Item": "Run ID", "Value": run.id},
            {"Item": "Fingerprint", "Value": manifest.fingerprint()},
            {"Item": "Generated", "Value": manifest.created_at},
        ]
    )
    if run.tournament and run.tournament.table:
        sheets["Model comparison"] = pd.DataFrame(run.tournament.table)
    if run.findings:
        sheets["Findings"] = pd.DataFrame(
            [
                {
                    "Title": f.title, "Detail": f.detail, "Evidence type": f.kind.value,
                    "Confidence": f.confidence.value, "Evidence": " | ".join(f.evidence),
                    "Caveats": " | ".join(f.caveats),
                }
                for f in run.findings
            ]
        )
    if run.recommendations:
        sheets["Recommendations"] = pd.DataFrame(
            [
                {
                    "Category": r.category, "Action": r.action, "Reason": r.reason,
                    "Confidence": r.confidence.value, "Expected impact": r.expected_impact,
                    "Evidence": " | ".join(r.evidence), "Caveats": " | ".join(r.caveats),
                    "Traceable to": " | ".join(r.traceable_to),
                }
                for r in run.recommendations
            ]
        )
    if run.profile:
        sheets["Variable profile"] = pd.DataFrame(
            [
                {
                    "Variable": name, "Type": col.semantic_type.value, "Dtype": col.dtype,
                    "Missing %": col.missing_pct, "Distinct": col.n_unique,
                    "Mean": col.mean, "Median": col.median, "Std": col.std,
                    "Min": col.minimum, "Max": col.maximum, "Skewness": col.skewness,
                    "Outliers %": col.outlier_pct,
                }
                for name, col in run.profile.columns.items()
            ]
        )
        if run.profile.quality_issues:
            sheets["Data quality"] = pd.DataFrame(
                [
                    {
                        "Severity": i.severity, "Code": i.code, "Message": i.message,
                        "Columns": ", ".join(i.columns), "Suggested action": i.suggested_action,
                    }
                    for i in run.profile.quality_issues
                ]
            )
    if run.explanation and run.explanation.importances:
        sheets["Feature importance"] = pd.DataFrame(
            [
                {"Rank": f.rank, "Variable": f.feature, "Importance": f.importance,
                 "Std": f.std, "Direction": f.direction, "Method": f.method}
                for f in run.explanation.importances
            ]
        )
    if run.decisions:
        sheets["Decision log"] = pd.DataFrame(
            [
                {
                    "Stage": d.stage, "Decision": d.decision, "Reason": d.reason,
                    "Evidence": " | ".join(d.evidence),
                    "Rejected": " | ".join(f"{r['option']}: {r['reason']}" for r in d.rejected),
                    "Confidence": d.confidence.value, "Overridden by user": d.overridden_by_user,
                }
                for d in run.decisions
            ]
        )
    if run.segmentation and getattr(run.segmentation, "clusters", None):
        sheets["Segments"] = pd.DataFrame(
            [
                {"Segment": c.name, "Label": c.label, "Size": c.size, "Share": c.share,
                 "Description": c.description}
                for c in run.segmentation.clusters
            ]
        )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, sheet in sheets.items():
            sheet.to_excel(writer, sheet_name=name[:31], index=False)
        _write_chart_sheet(writer, run, frame)
    return path


def _write_chart_sheet(writer: Any, run: Any, frame: pd.DataFrame | None) -> None:
    """A "Charts" sheet, each chart under its title and caption.

    Static images need a headless browser. Where there is none the sheet is still
    written with the titles, captions and the numbers behind each chart, so the
    workbook says everything the pictures would have said.
    """
    try:
        from dsai.reporting.figures import build_figures
    except Exception:
        return
    try:
        figures = build_figures(run, frame=frame)
    except Exception:
        return
    if not figures:
        return

    worksheet = writer.book.create_sheet("Charts")
    row = 1
    for figure in figures:
        worksheet.cell(row=row, column=1, value=figure.title).font = _bold_font()
        worksheet.cell(row=row + 1, column=1, value=figure.caption)
        row += 3
        png = to_png(figure.figure, width=900, height=figure.height)
        if png:
            try:
                from openpyxl.drawing.image import Image as XlImage

                image = XlImage(_io.BytesIO(png))
                image.width, image.height = 720, int(720 * figure.height / 900)
                worksheet.add_image(image, f"A{row}")
                row += int(image.height / 19) + 2
            except Exception:
                pass
        if figure.has_table:
            table = figure.table.head(30)
            for offset, column in enumerate(table.columns, start=1):
                worksheet.cell(row=row, column=offset, value=str(column)).font = _bold_font()
            for r, record in enumerate(table.itertuples(index=False), start=row + 1):
                for offset, value in enumerate(record, start=1):
                    worksheet.cell(row=r, column=offset,
                                   value=value if isinstance(value, (int, float, str)) else str(value))
            row += len(table) + 2
        row += 2
    worksheet.column_dimensions["A"].width = 34


def _bold_font():
    from openpyxl.styles import Font

    return Font(bold=True)


def to_json(run: Any, path: str | Path | None = None) -> str:
    from dsai.repro.provenance import build_manifest

    payload = {
        "run_id": run.id,
        "summary": run.summary(),
        "manifest": build_manifest(run).to_dict(),
        "findings": [f.to_dict() for f in run.findings],
        "recommendations": [r.to_dict() for r in run.recommendations],
        "decisions": [d.to_dict() for d in run.decisions],
        "trace": [e.to_dict() for e in run.trace.events],
        "tournament": run.tournament.to_dict() if run.tournament else None,
        "self_check": run.self_check.to_dict() if run.self_check else None,
        "profile": run.profile.to_dict() if run.profile else None,
    }
    text = json.dumps(payload, indent=2, default=str)
    if path:
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        Path(path).write_text(text, encoding="utf-8")
    return text


def to_pdf(run: Any, path: str | Path, audience: str = "both",
           frame: pd.DataFrame | None = None, methodology: bool = False) -> Path | None:
    """PDF via reportlab if installed; otherwise write HTML and say so.

    A browser's "print to PDF" on the HTML output gives a better result than a
    half-hearted PDF renderer, so that is the honest fallback rather than a
    silent failure.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError:
        # Not path.with_suffix(".html") — that collides with the HTML export
        # sitting in the same directory and silently replaces it.
        fallback = path.with_name(f"{path.stem}_print.html")
        fallback.write_text(to_html(run, audience), encoding="utf-8")
        raise RuntimeError(
            f"PDF export needs reportlab (pip install reportlab). "
            f"An HTML version was written to {fallback} instead — print that to PDF from your browser."
        )

    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image as PdfImage

    report = build_report(run, audience, frame=frame, include_methodology=methodology)
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=9.5, leading=13.5, spaceAfter=6)
    heading_style = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=13, spaceBefore=14, spaceAfter=6)
    caption_style = ParagraphStyle("Caption", parent=body_style, fontSize=8.5, leading=11.5,
                                   textColor="#5c6470", spaceBefore=3, spaceAfter=10)
    figure_title_style = ParagraphStyle("FigTitle", parent=styles["Heading3"], fontSize=10.5,
                                        spaceBefore=10, spaceAfter=4)

    document = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title=report.title,
    )
    flowables = [Paragraph(_html.escape(report.title), styles["Title"]),
                 Paragraph(f"Generated {report.generated_at}", body_style), Spacer(1, 12)]
    if report.executive_summary:
        flowables.append(Paragraph("Executive summary", heading_style))
        for paragraph in report.executive_summary.split("\n\n"):
            flowables.append(Paragraph(_pdf_inline(paragraph), body_style))
    usable_width = document.width
    for heading, content in report.sections:
        flowables.append(Paragraph(_html.escape(heading), heading_style))
        for paragraph in content.split("\n\n"):
            if paragraph.strip():
                flowables.append(Paragraph(_pdf_inline(paragraph), body_style))
        for figure in report.figures_for(heading):
            flowables.append(Paragraph(_html.escape(figure.title), figure_title_style))
            png = to_png(figure.figure, width=900, height=figure.height)
            if png:
                # Scale to the text column so a wide chart is never cropped.
                reader = ImageReader(_io.BytesIO(png))
                width_px, height_px = reader.getSize()
                flowables.append(PdfImage(_io.BytesIO(png), width=usable_width,
                                          height=usable_width * height_px / width_px))
            flowables.append(Paragraph(_pdf_inline(figure.caption), caption_style))
    document.build(flowables)
    return path


def _pdf_inline(text: str) -> str:
    import re

    escaped = _html.escape(text.replace("\n", "<br/>"))
    escaped = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"`([^`]+?)`", r"<font face='Courier'>\1</font>", escaped)
    return escaped


def to_python(run: Any, path: str | Path | None = None, data_path: str = "your_data.csv") -> str:
    from dsai.repro.codegen import generate_script

    script = generate_script(run, data_path)
    if path:
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        Path(path).write_text(script, encoding="utf-8")
    return script


def write_charts(run: Any, directory: str | Path,
                 frame: pd.DataFrame | None = None) -> dict[str, str]:
    """Write each chart as a PNG. Returns figure key -> path relative to *directory*.

    An empty result means this machine has no headless browser for static image
    export, not that the run had nothing to draw. Callers fall back to captions
    and tables.
    """
    from dsai.reporting.figures import build_figures

    directory = Path(directory)
    charts = directory / "charts"
    written: dict[str, str] = {}
    try:
        figures = build_figures(run, frame=frame)
    except Exception:
        return written
    for figure in figures:
        png = to_png(figure.figure, width=900, height=figure.height)
        if not png:
            continue
        charts.mkdir(parents=True, exist_ok=True)
        target = charts / f"{figure.key}.png"
        target.write_bytes(png)
        written[figure.key] = f"charts/{target.name}"
    return written


def export_audit_package(run: Any, directory: str | Path, data_path: str = "your_data.csv",
                         frame: pd.DataFrame | None = None) -> dict[str, str]:
    """Everything needed to check this analysis without the platform.

    One directory, one fingerprint, and a README naming what each file is for.
    The point is that a reviewer — an auditor, a colleague, the analyst
    themselves in six months — can reconstruct what was done and why without
    asking anyone.
    """
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    written = export_all(run, directory, data_path=data_path, frame=frame)

    stem = f"{run.dataset_name}_{run.id}"
    fingerprint = ""
    try:
        from dsai.repro.provenance import build_manifest, save_manifest

        manifest = build_manifest(run)
        fingerprint = manifest.fingerprint()
        target = directory / "dataset_manifest.json"
        save_manifest(manifest, target)
        written["manifest"] = str(target)
    except Exception as exc:
        written["manifest"] = f"failed: {exc}"

    for label, filename, builder in [
        ("methodology", "methodology.md",
         lambda: build_report(run, "technical", frame=frame, include_charts=False,
                              include_methodology=True).to_markdown()),
        ("model_card", "model_card.md", lambda: _model_card_markdown(run, frame)),
    ]:
        try:
            (directory / filename).write_text(builder(), encoding="utf-8")
            written[label] = str(directory / filename)
        except Exception as exc:
            written[label] = f"failed: {exc}"

    for label, filename, payload in [
        ("findings", "findings.json", [f.to_dict() for f in run.findings]),
        ("recommendations", "recommendations.json", [r.to_dict() for r in run.recommendations]),
        ("decision_log", "decision_log.json", [d.to_dict() for d in run.decisions]),
        ("evidence", "evidence.json",
         run.ledger.to_dict() if getattr(run, "ledger", None) is not None else {}),
        ("trust", "trust.json", _trust_payload(run)),
        ("contract", "data_contract.json",
         run.contract.to_dict() if getattr(run, "contract", None) is not None else {}),
    ]:
        try:
            (directory / filename).write_text(json.dumps(payload, indent=2, default=str),
                                              encoding="utf-8")
            written[label] = str(directory / filename)
        except Exception as exc:
            written[label] = f"failed: {exc}"

    try:
        (directory / "README.md").write_text(_audit_readme(run, stem, fingerprint, written),
                                             encoding="utf-8")
        written["readme"] = str(directory / "README.md")
    except Exception as exc:
        written["readme"] = f"failed: {exc}"
    return written


def _model_card_markdown(run: Any, frame: pd.DataFrame | None) -> str:
    from dsai.reporting.model_card import build_model_card

    return build_model_card(run, frame=frame).to_markdown()


def _trust_payload(run: Any) -> dict[str, Any]:
    trust = getattr(run, "trust", None)
    if trust is None:
        return {}
    return {
        "evidence_strength": trust.score,
        "band": trust.band,
        "verdict": trust.verdict,
        "assumption_debt": trust.assumption_debt,
        "debt_band": trust.debt_band,
        "supporting": [f.statement for f in trust.supporting],
        "reducing": [f.statement for f in trust.reducing],
        "known_unknowns": trust.unknowns,
        "debt_items": trust.debt_items,
        "blocking": trust.blocking,
        "note": "Evidence strength is a weighted summary of validation checks, not a "
                "statistical confidence level.",
    }


def _audit_readme(run: Any, stem: str, fingerprint: str, written: dict[str, str]) -> str:
    trust = getattr(run, "trust", None)
    lines = [
        f"# Analysis package — {run.dataset_name}", "",
        f"Run `{run.id}` · fingerprint `{fingerprint}` · generated {run.created_at}", "",
        "Everything in this directory describes one analysis. Two packages with the same "
        "fingerprint came from the same dataset, objective, preprocessing, validation strategy, "
        "seed and model, and should carry the same numbers.", "",
    ]
    if run.objective is not None:
        lines += [f"**Question.** {run.objective.task_type.value.replace('_', ' ')}"
                  + (f" on `{run.objective.target}`" if run.objective.target else ""), ""]
    if run.best is not None and run.plan is not None:
        lines += [f"**Selected model.** {run.best.model_name}, "
                  f"{run.plan.primary_metric} = "
                  f"{human_number(run.best.primary(run.plan.primary_metric))} "
                  "on this dataset under the validation strategy used.", ""]
    if trust is not None:
        lines += [f"**Evidence strength.** {trust.score}/100 ({trust.band}). "
                  f"Assumption debt {trust.assumption_debt}/100. "
                  "Neither is a statistical confidence level.", ""]

    lines += ["## What is in here", "", "| File | What it is |", "| --- | --- |"]
    catalogue = [
        (f"{stem}.html", "The report, self-contained, with interactive charts. Start here."),
        (f"{stem}.md", "The same report as Markdown."),
        (f"{stem}_business.md", "The management version — findings and actions, no methodology."),
        (f"{stem}_methodology.html", "The report with the full workings included."),
        ("methodology.md", "How it was run, the arithmetic behind every figure, and whether "
                           "anything went wrong producing it."),
        ("model_card.md", "What the model is for and — more importantly — what it is not for."),
        ("data_contract.json", "What the analysis required of its data. Check new data against "
                               "this before scoring."),
        ("evidence.json", "Every claim's identifier and what supports it."),
        ("trust.json", "Evidence strength, what supports and weakens it, and what this analysis "
                       "cannot answer at all."),
        ("findings.json", "The findings, with confidence and caveats."),
        ("recommendations.json", "The recommended actions and what each rests on."),
        ("decision_log.json", "Every automated decision, its reason, and the options rejected."),
        ("dataset_manifest.json", "The reproducibility fingerprint and its inputs."),
        (f"{stem}_reproduce.py", "Standalone Python that reproduces the analysis without this "
                                 "platform."),
        (f"{stem}.xlsx", "The tables as a workbook."),
        ("charts/", "Every chart as a PNG."),
    ]
    for filename, description in catalogue:
        lines.append(f"| `{filename}` | {description} |")

    failed = {k: v for k, v in written.items() if isinstance(v, str) and v.startswith(("failed", "not written"))}
    if failed:
        lines += ["", "## Not written", "",
                  "These could not be produced. The reason is given rather than the file being "
                  "quietly absent.", ""]
        lines += [f"- **{k}** — {v}" for k, v in failed.items()]

    lines += [
        "", "## How to read this", "",
        "1. `" + f"{stem}.html" + "` for the analysis as written.",
        "2. `trust.json` or the *Why trust this* section for whether it can carry a decision.",
        "3. `model_card.md` before using the model on anything.",
        "4. `methodology.md` to check the arithmetic.",
        "5. `" + f"{stem}_reproduce.py" + "` to run it yourself.",
        "", "Every relationship described in these files is an association. This is observational "
        "data; nothing here establishes that changing one thing would change another.",
    ]
    return "\n".join(lines)


def export_all(run: Any, directory: str | Path, data_path: str = "your_data.csv",
               frame: pd.DataFrame | None = None) -> dict[str, str]:
    """Write every available export format into one directory."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    stem = f"{run.dataset_name}_{run.id}"
    written: dict[str, str] = {}

    # Rendered once and referenced by the Markdown exports. The HTML report
    # carries its charts inside itself and does not need these.
    try:
        chart_files = write_charts(run, directory, frame=frame)
    except Exception:
        chart_files = {}
    if chart_files:
        written["charts"] = str(directory / "charts")

    for label, filename, writer in [
        ("markdown", f"{stem}.md",
         lambda p: p.write_text(to_markdown(run, frame=frame, chart_files=chart_files), encoding="utf-8")),
        ("html", f"{stem}.html", lambda p: p.write_text(to_html(run, frame=frame), encoding="utf-8")),
        ("business_summary", f"{stem}_business.md",
         lambda p: p.write_text(to_markdown(run, "business", frame=frame, chart_files=chart_files),
                                encoding="utf-8")),
        # The full-method version is written alongside rather than instead: most
        # readers do not want it, and the one who does should not have to re-run
        # the export to get it.
        ("methodology", f"{stem}_methodology.html",
         lambda p: p.write_text(to_html(run, "technical", frame=frame, methodology=True),
                                encoding="utf-8")),
        ("json", f"{stem}.json", lambda p: to_json(run, p)),
        ("python", f"{stem}_reproduce.py", lambda p: to_python(run, p, data_path)),
        ("excel", f"{stem}.xlsx", lambda p: to_excel(run, p, frame=frame)),
    ]:
        target = directory / filename
        try:
            writer(target)
            written[label] = str(target)
        except Exception as exc:
            written[label] = f"failed: {exc}"

    try:
        written["pdf"] = str(to_pdf(run, directory / f"{stem}.pdf", frame=frame))
    except Exception as exc:
        written["pdf"] = f"not written: {exc}"
    return written
